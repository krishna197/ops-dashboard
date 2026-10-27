#!/usr/bin/env python3
"""Operations Dashboard - reads an Excel workbook and serves a live web dashboard.

    python app.py            start the dashboard (settings in config.ini)
    python app.py --check    show what was found in your Excel (use after changing config)
"""
import configparser, io, json, os, re, socket, sys, threading, webbrowser
from datetime import datetime, timezone
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is missing. Run:  pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
VERTICALS = ["HBC", "FnN", "PHC", "BEV", "GTX"]

# ---------------------------------------------------------------- config
CP = configparser.ConfigParser(interpolation=None, inline_comment_prefixes=(";",))
if not CP.read(os.path.join(HERE, "config.ini"), encoding="utf-8"):
    sys.exit("config.ini not found next to app.py")
A = CP["app"]


def excel_path():
    """DASH_EXCEL env var (for testing) overrides excel_file from config.ini."""
    return os.path.join(HERE, os.environ.get("DASH_EXCEL") or A.get("excel_file", "Data.xlsx"))


def split(v):
    return [x.strip() for x in v.split(",") if x.strip()]


def make_cfg():
    groups = []
    for name in CP.sections():
        if not name.startswith("group:"):
            continue
        s, charts = CP[name], []
        for k in sorted(k for k in s if k.startswith("chart")):
            p = [x.strip() for x in s[k].split("=>")]
            if len(p) >= 4:
                charts.append({"t": p[0], "s": p[1].lower() == "stacked", "r": split(p[2]), "c": split(p[3])})
        gid = name.split(":", 1)[1].strip()
        groups.append({"id": gid, "title": s.get("title", gid), "icon": s.get("icon", "•"),
                       "section": s.get("section", "Operations"), "match": s.get("match", "$^"), "charts": charts})
    for pat in [g["match"] for g in groups] + [CP.get("rules", "lower_is_better", fallback="$^")]:
        try:
            re.compile(pat)
        except re.error as e:
            sys.exit(f"config.ini: bad regex '{pat}' ({e})")
    r = CP["rules"] if CP.has_section("rules") else {}
    return {"title": A.get("title", "Operations Dashboard"), "theme": A.get("theme", "dark"),
            "refresh": A.getint("refresh_seconds", 30), "groups": groups,
            "headline": split(CP.get("overview", "headline", fallback="")),
            "lower": r.get("lower_is_better", "$^"), "good": float(r.get("good_score", 90)), "warn": float(r.get("warn_score", 70))}


CFG = make_cfg()

# ---------------------------------------------------------------- excel -> data
def num(v):
    if v is None or isinstance(v, bool):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def mon(v):
    if isinstance(v, datetime):
        return MONTHS[v.month - 1]
    return str(v or "").strip()[:3]


def grid(rows):
    """Metric x month (or metric x vertical) table. Header row = first row holding month/vertical names."""
    ok = lambda c: c in MONTHS or c in VERTICALS
    h = next((i for i, r in enumerate(rows) if any(ok(c) for c in r)), None)
    if h is None:
        return {"cols": [], "rows": []}
    c0 = next(i for i, c in enumerate(rows[h]) if ok(c))
    cols, c = [], c0
    while c < len(rows[h]) and isinstance(rows[h][c], str) and rows[h][c] != "Total":
        cols.append(rows[h][c]); c += 1
    out = []
    for r in rows[h + 1:]:
        r = list(r) + [None] * (c0 + len(cols))
        label = r[c0 - 1] if c0 > 0 else None
        if isinstance(label, str) and label.strip():
            out.append([label.strip(), [num(r[c0 + k]) for k in range(len(cols))]])
    return {"cols": cols, "rows": out}


def standardization(rows):
    names = ["Month Name", "Database Set-up", "Out of System Tool", "Validation", "Total Score"]
    h = next((i for i, r in enumerate(rows) if "Month Name" in r), None)
    if h is None:
        return []
    ix = [rows[h].index(n) if n in rows[h] else None for n in names]
    out = []
    for r in rows[h + 1:]:
        r = list(r) + [None] * 30
        if ix[0] is not None and r[ix[0]] in MONTHS:
            out.append([r[ix[0]]] + [num(r[i]) if i is not None else None for i in ix[1:]])
    return out


def kpi(rows):
    if not rows:
        return []
    ix = {n: i for i, n in enumerate(rows[0]) if n}
    g = lambda r, n: r[ix[n]] if n in ix and ix[n] < len(r) else None
    out = []
    for r in rows[1:]:
        if not g(r, "KPI Name"):
            continue
        c = g(r, "Compliant?")
        out.append([mon(g(r, "Month")), g(r, "Output Process Name") or "", g(r, "Area") or "", g(r, "KPI Name"),
                    g(r, "Actual Value"), g(r, "Target Value"), c if c in ("Y", "N") else "NA"])
    return out


def build(path):
    with open(path, "rb") as f:                       # read bytes: works even while Excel has it open
        raw = f.read()
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    warn, S = [], CP["sheets"]

    def sheet(key):
        want = S.get(key, "").lower()
        name = next((n for n in wb.sheetnames if want and want in n.lower()), None)
        if not name:
            warn.append(f"Sheet containing '{want}' not found")
            return []
        return [list(r) for r in wb[name].iter_rows(values_only=True)]

    data = {"g": {"ov": grid(sheet("overall")), "db": grid(sheet("database")), "vt": grid(sheet("vertical"))},
            "sd": standardization(sheet("standardization")), "kp": kpi(sheet("kpi"))}
    wb.close()
    return data, warn


# ---------------------------------------------------------------- cache (re-reads only when the file is saved)
_lock, _st = threading.Lock(), {"mtime": None, "payload": None}


def payload():
    path = excel_path()
    with _lock:
        try:
            mt = os.stat(path).st_mtime
        except OSError:
            return {"version": 0, "cfg": CFG, "error": f"Excel file not found: {path}  (put it there or change excel_file in config.ini)"}
        if mt != _st["mtime"]:
            try:
                data, warn = build(path)
                _st.update(mtime=mt, payload={"version": int(mt * 1000), "cfg": CFG, "error": "", "data": data, "warnings": warn,
                                              "file": os.path.basename(path),
                                              "updated": datetime.fromtimestamp(mt, timezone.utc).isoformat()})
            except Exception as e:                      # e.g. file half-saved: keep showing last good data
                msg = f"Could not read {os.path.basename(path)} ({e.__class__.__name__}). Save the file again."
                if _st["payload"]:
                    return dict(_st["payload"], warnings=_st["payload"]["warnings"] + [msg])
                return {"version": -1, "cfg": CFG, "error": msg}
        return _st["payload"]


# ---------------------------------------------------------------- web server
class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *a, **k):
        super().__init__(*a, directory=os.path.join(HERE, "static"), **k)

    def do_GET(self):
        route = self.path.split("?")[0]
        if route == "/api/data":
            body = json.dumps(payload(), default=str).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if route == "/":
            self.path = "/index.html"
        super().do_GET()

    def end_headers(self):
        self.send_header("Cache-Control", "no-store")
        super().end_headers()

    def log_message(self, *a):
        pass


def lan_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM); s.connect(("10.255.255.255", 1)); ip = s.getsockname()[0]; s.close()
        return ip
    except Exception:
        return "localhost"


def check():
    p = payload()
    print(f"\nExcel : {os.path.basename(excel_path())}")
    if p.get("error"):
        print("ERROR :", p["error"]); return
    d = p["data"]
    print(f"Saved : {p['updated']}")
    print(f"Found : {len(d['g']['ov']['rows'])} overall metrics, {len(d['g']['db']['rows'])} database rows, "
          f"{len(d['g']['vt']['rows'])} vertical rows, {len(d['sd'])} score months, {len(d['kp'])} KPI rows")
    for w in p["warnings"]:
        print("WARN  :", w)
    used = set()
    print("\nTabs:")
    for g in CFG["groups"]:
        names = [r[0] for r in d["g"]["ov"]["rows"] if re.search(g["match"], r[0], re.I)]
        used.update(names)
        print(f"  {g['title']:<18} {len(names)} metrics")
    skip = re.compile(r"standardization|database setup|out of system|validation", re.I)
    left = [r[0] for r in d["g"]["ov"]["rows"] if r[0] not in used and not skip.search(r[0])]
    if left:
        print("\nMetrics not in any tab (add a [group:...] in config.ini if you want them):")
        for n in left:
            print("  -", n)
    print()


def main():
    if "--check" in sys.argv:
        return check()
    host, port = A.get("host", "0.0.0.0"), A.getint("port", 8000)
    try:
        srv = ThreadingHTTPServer((host, port), Handler)
    except OSError:
        sys.exit(f"Port {port} is already in use. Change 'port' in config.ini.")
    print(f"\n  {CFG['title']} is running\n")
    print(f"  You       : http://localhost:{port}")
    if host != "localhost":
        print(f"  Teammates : http://{lan_ip()}:{port}   (or http://{socket.gethostname()}:{port})")
    print(f"\n  Edit and SAVE {A.get('excel_file')} - dashboards refresh by themselves.")
    print("  Keep this window open. Press Ctrl+C to stop.\n")
    if A.getboolean("open_browser", True):
        threading.Timer(1, lambda: webbrowser.open(f"http://localhost:{port}")).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
