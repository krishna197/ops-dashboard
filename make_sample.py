"""Creates Sample_Data.xlsx: FAKE numbers in the exact layout the dashboard expects.
Use it as a blank template for teammates, or to preview the dashboard:  set excel_file = Sample_Data.xlsx"""
import random
from openpyxl import Workbook

random.seed(7)
M = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
V = ["HBC", "FnN", "PHC", "BEV", "GTX"]
N = 9  # months filled (Jan-Sep)
BASE = {"Clients": 130, "Databases": 450, "DB Health - Red": 6, "DB Health - Amber": 24, "DB Health - Green": 420, "New UPCs": 3200,
        "Dropped UPCs": 900, "Value Change": 52000, "Unique Value Change": 18000, "Not Included": 400, "Missed Observation in ARS": 12,
        "Change Requests Received": 60, "Client Inquiries Received": 240, "Total Cycle Time": 52, "ION Cycle Time (NOVA)": 31,
        "Say Do Ration (NOVA)": 94, "Quality Escape": 4, "Mid Cyle / Mid week Production": 18, "Standardization RMS": 86,
        "Database Setup": 70, "Out of System Tool": 95, "Validation": 85, "New Build Received (Build Team)": 5, "New Build to BAU (BAU)": 4,
        "Team Size": 62, "Available Size": 55, "Attrition (Relieved and not in notice period)": 2, "New Joinees": 3, "Saviom Compliance": 96}
LABELS = list(BASE) + ["Change Requests Implemented", "Client Inquiries Responded"]
LABELS.sort(key=lambda x: ["Clients", "Databases", "DB Health - Red", "DB Health - Amber", "DB Health - Green"].index(x) if x in ["Clients", "Databases", "DB Health - Red", "DB Health - Amber", "DB Health - Green"] else 99)
series = {k: [round(v * random.uniform(.88, 1.12), 1) for _ in range(N)] for k, v in BASE.items()}
series["Change Requests Implemented"] = [round(x * random.uniform(.8, 1)) for x in series["Change Requests Received"]]
series["Client Inquiries Responded"] = [round(x * random.uniform(.85, 1)) for x in series["Client Inquiries Received"]]

wb = Workbook()
ws = wb.active; ws.title = "1. Overall"
for i, m in enumerate(M): ws.cell(2, 3 + i, m)
for r, k in enumerate(LABELS):
    ws.cell(3 + r, 2, k)
    for i, v in enumerate(series[k]): ws.cell(3 + r, 3 + i, v)

ws = wb.create_sheet("2. Database Overview")
for i, m in enumerate(M): ws.cell(3, 3 + i, m)
for r, (v, kind) in enumerate([(v, k) for v in V for k in ("Clients", "Database")]):
    ws.cell(4 + r, 2, f"{v} - {kind}")
    for i in range(12): ws.cell(4 + r, 3 + i, random.randint(15, 40) if kind == "Clients" else random.randint(60, 130))

ws = wb.create_sheet("Vertical Overview")
for i, v in enumerate(V): ws.cell(3, 3 + i, v)
for r, k in enumerate(["Clients", "Database", "New UPCs", "Dropped UPCs", "Quality Escape", "Team Size"]):
    ws.cell(4 + r, 2, k)
    for i in range(5): ws.cell(4 + r, 3 + i, random.randint(10, 120) if k in ("Clients", "Database", "Team Size") else random.randint(0, 900))

ws = wb.create_sheet("Standardization Overview")
for i, h in enumerate(["Month #", "Month Name", "Database Set-up", "Out of System Tool", "Validation", "Total Score"]): ws.cell(2, 2 + i, h)
for i, m in enumerate(M):
    ws.cell(3 + i, 2, i + 1); ws.cell(3 + i, 3, m)
    if i < N:
        a, b, c = round(random.uniform(55, 90), 2), round(random.uniform(93, 97), 2), round(random.uniform(70, 97), 2)
        for j, v in enumerate((a, b, c, round((a + b + c) / 3, 2))): ws.cell(3 + i, 4 + j, v)

ws = wb.create_sheet("Standardization KPI Trend")
for i, h in enumerate(["Month", "Country", "Service", "Function", "Output Process Name", "SOP ID", "Area", "KPI Name", "Actual Value", "Target Value", "Compliant?"]): ws.cell(1, 1 + i, h)
KP = [("Database Set-up", "Core KPIs", f"Sample KPI {i}") for i in range(1, 9)] + [("Validation", "Essential KPIs", f"Sample check {i}") for i in range(1, 7)] + [("Production", "Essential KPIs", f"Sample run {i}") for i in range(1, 4)]
row = 2
for mi in range(N):
    for proc, area, name in KP:
        ok = random.random() > .25
        for j, v in enumerate((f"{M[mi]} 2026", "USA", "RMS", "Output", proc, None, area, name, round(random.uniform(0, 10), 2), "<=10", random.choice(["Y", "Y", "Y", "N", "NA"]) if ok else "N")):
            ws.cell(row, 1 + j, v)
        row += 1
wb.save("Sample_Data.xlsx"); print("Created Sample_Data.xlsx (fake demo numbers, same layout as your workbook)")
